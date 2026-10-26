"""
Pendencias de cadastro de transferencias por UG (subsidio ao envio de processos SEI).

Compara o SIGGo (MIL2026.TRANSFERENCIA) com o TransfereGov (GDF, instrumentos em fase ativa):
  * "Sem cadastro identificado": instrumento cujo identificador nao esta no
    NUTRANSFSIAFI de nenhuma transferencia e sem correspondencia segura no SIGGo
    (UG pelo CNPJ do beneficiario/executor);
  * "Cadastro inconsistente": transferencia do SIGGo identificada (Nº Original ou
    conta exata) com campos divergentes do TransfereGov (IN 13/2026).

Usa a extracao ja feita (rode antes extrair_transferegov_df.py, reparar_extracao_transferegov.py
e extrair_siggo_transferencia.py). Saida: pendencias_por_ug_AAAAMMDD.xlsx.
"""

import re
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import painel_transferegov_base as g
import gerar_relatorio_pendencias_transferegov as rel

COLUNAS = ["UG", "Situação", "Nº Transferência (SIGGo)", "Identificador (TransfereGov)", "Modalidade", "Concedente",
           "Beneficiário / Executor", "Início da Vigência", "Fim da Vigência", "Valor Global (R$)",
           "Inconsistências (SIGGo × TransfereGov) e providência"]
GENERICOS = ("distrito federal", "governo do distrito federal", "estado do distrito federal")


def titulo(x):
    t = str(x).strip().title() if pd.notna(x) and str(x).strip() else ""
    return {"Convenio": "Convênio"}.get(t, t)


def dig(x):
    return re.sub(r"\D", "", str(x if pd.notna(x) else "")).lstrip("0")


def dt(x, dayfirst=True):
    d = pd.to_datetime(x, dayfirst=dayfirst, errors="coerce")
    return d.to_pydatetime() if pd.notna(d) else None


def dmy(d):
    return d.strftime("%d/%m/%Y") if d else "—"


def brl(v):
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def detalhes_siconv():
    prop = g.carregar("df_siconv_proposta.csv", dtype=str)
    conv = g.carregar("df_siconv_convenio.csv", dtype=str)
    m = conv.merge(prop, on="ID_PROPOSTA", how="left").drop_duplicates("NR_CONVENIO")
    out = {}
    for _, r in m.iterrows():
        ini = r["DIA_INIC_VIGENC_CONV"] if pd.notna(r["DIA_INIC_VIGENC_CONV"]) else r["DIA_INIC_VIGENCIA_PROPOSTA"]
        fim = r["DIA_FIM_VIGENC_CONV"] if pd.notna(r["DIA_FIM_VIGENC_CONV"]) else r["DIA_FIM_VIGENCIA_PROPOSTA"]
        esp = {"CONVENIO": "1", "CONTRATO DE REPASSE": "13"}.get(str(r["MODALIDADE"]).strip().upper())
        out[g.chave_id(r["NR_CONVENIO"])] = dict(mod=titulo(r["MODALIDADE"]), benef=titulo(r["NM_PROPONENTE"]), ini=dt(ini), fim=dt(fim),
                                                 cnpj=dig(r["IDENTIF_PROPONENTE"]), conta=dig(r["CD_CONTA"]), esp=esp)
    return out


def detalhes_parcerias():
    pp = g.carregar("df_parcerias_proposta.csv", dtype=str)
    par = g.carregar("df_parcerias_parceria.csv", dtype=str)
    prog = g.carregar("df_parcerias_programa.csv", dtype=str)
    m = par.merge(pp, on="id_proposta", how="left").merge(prog[["id_programa", "tp_instrumento"]], on="id_programa", how="left")
    out = {}
    for _, r in m.iterrows():
        e = pd.notna(r["nm_executor"])
        out[g.chave_id(r["cd_parceria"])] = dict(mod=titulo(r["tp_instrumento"]) or "Gestão de Parcerias",
                                                 benef=titulo(r["nm_executor"] if e else r["nm_ente_recebedor"]), ini=None, fim=None,
                                                 cnpj=dig(r["cnpj_executor"] if e else r["cnpj_ente_recebedor"]), conta=None, esp="1")
    return out


def detalhes_especiais():
    pa = g.carregar("df_especiais_plano_acao.csv", dtype=str)
    ex = g.carregar("df_especiais_executor.csv", dtype=str).drop_duplicates("id_plano_acao")
    pt = g.carregar("df_especiais_plano_trabalho.csv", dtype=str).drop_duplicates("id_plano_acao")
    bn = g.carregar("df_especiais_beneficiario.csv", dtype=str)
    m = pa.merge(ex[["id_plano_acao", "nome_executor", "cnpj_executor"]], on="id_plano_acao", how="left") \
          .merge(pt[["id_plano_acao", "data_inicio_execucao_plano_trabalho", "data_fim_execucao_plano_trabalho"]], on="id_plano_acao", how="left") \
          .merge(bn[["id_beneficiario", "nome_beneficiario", "cnpj_beneficiario"]], on="id_beneficiario", how="left")
    out = {}
    for _, r in m.iterrows():
        e = pd.notna(r["nome_executor"])
        out[g.chave_id(r["codigo_plano_acao"])] = dict(mod="Transferência Especial", benef=titulo(r["nome_executor"] if e else r["nome_beneficiario"]),
                                                       ini=dt(r["data_inicio_execucao_plano_trabalho"], False), fim=dt(r["data_fim_execucao_plano_trabalho"], False),
                                                       cnpj=dig(r["cnpj_executor"] if e else r["cnpj_beneficiario"]), conta=None, esp="10")
    return out


def detalhes_faf():
    pa = g.carregar("df_fundoafundo_plano_acao.csv", dtype=str).drop_duplicates("codigo_plano_acao")
    out = {}
    for _, r in pa.iterrows():
        v = pd.notna(r["nome_fundo_vinculado_plano_acao"])
        out[g.chave_id(r["codigo_plano_acao"])] = dict(
            mod="Fundo a Fundo", benef=titulo(r["nome_fundo_vinculado_plano_acao"] if v else r["nome_ente_recebedor_plano_acao"]),
            ini=dt(r["data_inicio_vigencia_plano_acao"], False), fim=dt(r["data_fim_vigencia_plano_acao"], False),
            cnpj=dig(r["cnpj_fundo_vinculado_plano_acao"] if v else r["cnpj_ente_recebedor_plano_acao"]), conta=None, esp="12")
    return out


def id_esperado(fonte, ident):
    """Valor esperado em Nº TransfereGov (IN 13, art. 6º): so digitos; vazio se o codigo e alfanumerico (SICONV 7AA...)."""
    ident = str(ident).split(".")[0] if fonte == "Gestão de Parcerias" else str(ident)
    if fonte == "Discricionárias e Legais":
        s = re.sub(r"[^0-9A-Za-z]", "", ident)
        return s if s.isdigit() else ""
    return re.sub(r"\D", "", ident)


def inconsistencias(r, p, d):
    itens = []
    esperado = id_esperado(p["Fonte"], p["Identificador"])
    tg = str(r["NUTRANSFSIAFI"]).strip()
    if esperado and tg != esperado:
        itens.append(f"Nº TransfereGov: registrado {'em branco' if not tg else repr(tg)} → informar {esperado} (somente dígitos; art. 5º, §§1º e 2º)")
    elif not esperado and tg not in ("", "None"):
        itens.append(f"Nº TransfereGov: registrado {tg!r} → código alfanumérico; deixar em branco (art. 6º, I, \"b\")")
    no = str(r["NUORIGINAL"]).strip()
    if p["Fonte"] in ("Discricionárias e Legais", "Gestão de Parcerias") and esperado and esperado not in re.sub(r"\D", "", no):
        itens.append(f"Número Original: registrado {no!r} → deve conter o identificador {esperado} (art. 6º)")
    if d.get("esp") and str(r["INESPECIE"]).strip().lstrip("0") != d["esp"]:
        itens.append(f"Espécie de Transferência: registrada {str(r['INESPECIE']).strip()} → {d['esp']} ({d['mod']}; art. 5º, §6º)")
    va = float(g.num(pd.Series([r["VATRANSFERENCIA"]])).iloc[0])
    vg = float(p["Valor Global (R$)"])
    if p["Fonte"] != "Fundo a Fundo" and vg > 0 and abs(va - vg) >= 0.01:  # FaF: o SIGGo registra por repasse, nao o total do plano
        itens.append(f"Valor Transferência: registrado {brl(va)} → {brl(vg)} (valor global atualizado; art. 5º, §8º)")
    for nome, col, ref in (("Data Início Vigência", "DAINIVIGENCIA", d.get("ini")), ("Data Fim Vigência", "DAFIMVIGENCIA", d.get("fim"))):
        a = dt(r[col], False)
        if ref and a and a.date() != ref.date():
            itens.append(f"{nome}: registrada {dmy(a)} → {dmy(ref)} (art. 5º, §11)")
    if d.get("conta"):
        sg, tg_ = dig(r["NUCONTA"]), d["conta"]
        # mesma conta pode vir sem o codigo da operacao (ex.: 000710687 x 0060710687): compara pelo final
        if not (sg == tg_ or (len(sg) >= 6 and tg_.endswith(sg)) or (len(tg_) >= 6 and sg.endswith(tg_))):
            itens.append(f"Conta Bancária: registrada {str(r['NUCONTA']).strip() or 'em branco'} → {d['conta']} (conta do instrumento no TransfereGov, a comprovar; art. 5º, §10)")
    return itens


def main():
    pend, todos, _ = rel.montar_universo()

    siggo = pd.read_csv(g.PASTA / "df_siggo_transferencia.csv", dtype=str, low_memory=False).fillna("")
    siggo["NT"] = siggo["NUTRANSFERENCIA"].str.strip()
    siggo = siggo.drop_duplicates("NT").set_index("NT")
    ugs = pd.read_csv(g.PASTA / "unidadegestora_bruto.csv", dtype=str).fillna("")
    nome_ug = {r.COUG.strip(): r.NOUG.strip() for r in ugs.itertuples()}
    ug_por_cnpj = {dig(r.NUCGC): r.COUG.strip() for r in ugs.itertuples() if dig(r.NUCGC)}

    det = {"Discricionárias e Legais": detalhes_siconv(), "Gestão de Parcerias": detalhes_parcerias(),
           "Transferências Especiais": detalhes_especiais(), "Fundo a Fundo": detalhes_faf()}

    def rotulo_ug(ug):
        return f"{ug} – {nome_ug.get(ug, '')}".strip(" –") if ug else "UG não identificada"

    def benef_detalhado(benef, ug, fallback):
        b = benef or titulo(fallback)
        if re.sub(r"[^a-z ]", "", b.lower()).strip() in GENERICOS and ug:
            return titulo(nome_ug.get(ug, b)) + " (UG no SIGGo)"
        return b

    linhas = []
    for _, p in pend.iterrows():
        d = det[p["Fonte"]].get(g.chave_id(p["Identificador"]), {})
        cands = [(nt, cam) for nt, cam in re.findall(r"(\d+) \(camada (\d)\)", str(p["Candidato no SIGGo (Nº Transf. / camada)"]))
                 if cam in ("2", "3") and nt in siggo.index]
        mod = d.get("mod") or p["Fonte"]
        conc = titulo(p["Concedente"])
        if cands:
            for nt, _cam in cands:
                s = siggo.loc[nt]
                ben, con, reg = s["COBENEFICIADO"].strip(), s["COCONCENTE"].strip(), s["COUG"].strip()
                if reg not in ("", "0", "130101"):
                    ug = reg
                elif re.fullmatch(r"\d{6}-\d{5}", ben):
                    ug = ben.split("-")[0]
                elif dig(ben) in ug_por_cnpj:
                    ug = ug_por_cnpj[dig(ben)]
                elif re.fullmatch(r"\d{6}-\d{5}", con):
                    ug = con.split("-")[0]
                else:
                    ug = reg
                inc = inconsistencias(s, p, d)
                linhas.append([rotulo_ug(ug), "Cadastro inconsistente", int(nt), p["Identificador"], mod, conc,
                               benef_detalhado(d.get("benef"), ug, p["Beneficiário / Órgão executor"]), d.get("ini"), d.get("fim"),
                               float(p["Valor Global (R$)"]), "\n".join(inc) or "Identificação a confirmar; demais campos sem divergência apurada"])
        else:
            ug = ug_por_cnpj.get(d.get("cnpj", ""), "")
            linhas.append([rotulo_ug(ug), "Sem cadastro identificado", None, p["Identificador"], mod, conc,
                           benef_detalhado(d.get("benef"), ug, p["Beneficiário / Órgão executor"]), d.get("ini"), d.get("fim"),
                           float(p["Valor Global (R$)"]),
                           "Instrumento não localizado no SIGGo: cadastrar a transferência (AFC) com Nº TransfereGov "
                           f"{id_esperado(p['Fonte'], p['Identificador']) or '(em branco; código alfanumérico)'}, conforme IN nº 13/2026"])
    # transferencias ja identificadas pelo Nº TransfereGov (camada 1), mas com outros campos divergentes
    for _, p in todos[(todos["Localizado no SIGGo"] == "Sim") & (todos["Em fase ativa"] == "Sim")].iterrows():
        d = det[p["Fonte"]].get(g.chave_id(p["Identificador"]), {})
        for nt in [x.strip() for x in str(p["Nº Transf. SIGGo"]).split(",") if x.strip() in siggo.index]:
            sg = siggo.loc[nt]
            inc = inconsistencias(sg, p, d)
            if not inc:
                continue
            ben, con, reg = sg["COBENEFICIADO"].strip(), sg["COCONCENTE"].strip(), sg["COUG"].strip()
            if reg not in ("", "0", "130101"):
                ug = reg
            elif re.fullmatch(r"\d{6}-\d{5}", ben):
                ug = ben.split("-")[0]
            elif dig(ben) in ug_por_cnpj:
                ug = ug_por_cnpj[dig(ben)]
            else:
                ug = reg
            linhas.append([rotulo_ug(ug), "Cadastro inconsistente", int(nt), p["Identificador"], d.get("mod") or p["Fonte"], titulo(p["Concedente"]),
                           benef_detalhado(d.get("benef"), ug, p["Beneficiário / Órgão executor"]), d.get("ini"), d.get("fim"),
                           float(p["Valor Global (R$)"]), chr(10).join(inc)])
    df = pd.DataFrame(linhas, columns=COLUNAS).drop_duplicates(subset=COLUNAS[:4])
    df["_o"] = df["Situação"].map({"Cadastro inconsistente": 0, "Sem cadastro identificado": 1})
    df = df.sort_values(["UG", "_o", "Nº Transferência (SIGGo)", "Identificador (TransfereGov)"]).drop(columns="_o")

    resumo = df.groupby("UG").apply(lambda x: pd.Series({
        "Cadastro inconsistente": int((x["Situação"] == "Cadastro inconsistente").sum()),
        "Sem cadastro identificado": int((x["Situação"] == "Sem cadastro identificado").sum()),
        "Total": len(x), "Valor Global (R$)": float(x["Valor Global (R$)"].sum())}), include_groups=False).reset_index()
    resumo = resumo.sort_values("Total", ascending=False)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Resumo por UG"
    ws = wb.create_sheet("Detalhamento")

    def cab(w, cols):
        w.append(cols)
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0D1B3E")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        w.freeze_panes = "A2"

    cab(ws0, list(resumo.columns))
    for r in resumo.itertuples(index=False):
        ws0.append(list(r))
    n = ws0.max_row
    ws0.append(["TOTAL"] + [f"=SUM({get_column_letter(i)}2:{get_column_letter(i)}{n})" for i in range(2, 6)])
    for c in ws0[ws0.max_row]:
        c.font = Font(bold=True)
    for row in ws0.iter_rows(min_row=2):
        row[4].number_format = "#,##0.00"
    for i, w in enumerate([58, 18, 22, 10, 22], 1):
        ws0.column_dimensions[get_column_letter(i)].width = w

    cab(ws, COLUNAS)
    for r in df.itertuples(index=False):
        ws.append(list(r))
    for i, w in enumerate([44, 22, 14, 24, 24, 32, 42, 13, 13, 18, 90], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[7].number_format = row[8].number_format = "dd/mm/yyyy"
        row[9].number_format = "#,##0.00"
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (5, 6, 7, 11))
        row[1].fill = PatternFill("solid", fgColor="FFC7CE" if row[1].value == "Sem cadastro identificado" else "FFEB9C")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ws.max_row}"

    hoje = datetime.now()
    saida = g.PASTA / f"pendencias_por_ug_{hoje:%Y%m%d}.xlsx"
    try:
        wb.save(saida)
    except PermissionError:
        saida = saida.with_name(f"{saida.stem}_{hoje:%H%M}.xlsx")
        wb.save(saida)
    print(saida)
    print(len(df), "linhas;", df["UG"].nunique(), "UGs")
    print(df["Situação"].value_counts().to_string())
    print(resumo.head(30).to_string())


if __name__ == "__main__":
    main()
