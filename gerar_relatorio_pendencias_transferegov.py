"""
Relatorio (Excel) de PENDENCIAS de cadastro no SIGGo: instrumentos que constam
nas APIs/CSVs do TransfereGov (somente GDF) e cujo identificador NAO foi
cadastrado no campo NUTRANSFSIAFI de nenhum registro do SIGGo (camada 1 do
cruzamento - unica considerada identificacao; as demais camadas viram apenas
"candidato" para conferencia).

Abas:
  Resumo        - totais por fonte
  Pendencias    - nao localizados, em fase ativa, COM e SEM execucao financeira
  Todos (APIs)  - todos os instrumentos das APIs; verde = localizado no SIGGo,
                  vermelho = nao localizado

Antes de gerar, rode extrair_siggo_transferencia.py e extrair_transferegov_df.py
(e reparar_extracao_transferegov.py se a API deu erro 502).
Saida: relatorio_pendencias_transferegov_AAAAMMDD.xlsx (nesta pasta).
"""

import re
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import gerar_painel_transferegov_df as g

COLUNAS = ["Fonte", "Identificador", "Concedente", "Beneficiário / Órgão executor", "Objeto",
           "Data de Celebração", "Data da Proposta", "Valor Global (R$)", "Valor Repassado (R$)", "Situação", "Nº da Emenda",
           "Execução financeira", "Candidato no SIGGo (Nº Transf. / camada)"]
COLUNAS_PEND = ["Fonte", "Identificador", "Concedente", "Beneficiário / Órgão executor",
                "Data de Celebração", "Data da Proposta", "Valor Global (R$)", "Valor Repassado (R$)", "Situação", "Nº da Emenda",
                "Execução financeira"]
COLUNAS_TODOS = COLUNAS + ["Localizado no SIGGo", "NUTRANSFSIAFI no SIGGo", "Nº Transf. SIGGo", "Em fase ativa"]

# Situacoes que tiram o instrumento do universo de pendencias (encerrados sem
# efeito ou ainda em elaboracao): anulado, cancelado, rescindido, impedido,
# rejeitado, em elaboracao.
INATIVAS = ("ANULAD", "CANCELAD", "RESCINDID", "IMPEDID", "REJEITAD", "ELABORACAO", "ELABORAÇÃO")

VERDE = PatternFill("solid", fgColor="C6EFCE")
VERMELHO = PatternFill("solid", fgColor="FFC7CE")


def em_fase_ativa(situacao, executado, tem_data):
    sit = str(situacao or "").strip().upper()
    if not sit or sit == "NAN":  # situacao nao informada: so entra se ha execucao ou data
        return executado > 0 or tem_data
    return not any(m in sit for m in INATIVAS)


def objeto_siconv():
    prop = g.carregar("df_siconv_proposta.csv", dtype=str)
    conv = g.carregar("df_siconv_convenio.csv", dtype=str)
    m = conv.merge(prop[["ID_PROPOSTA", "OBJETO_PROPOSTA"]], on="ID_PROPOSTA", how="left")
    return m.drop_duplicates("NR_CONVENIO").set_index(m["NR_CONVENIO"].astype(str).str.strip())["OBJETO_PROPOSTA"]


def datas_proposta():
    """Data da proposta por identificador, para os instrumentos sem data de celebracao:
    Gestao de Parcerias (nenhuma parceria do GDF traz data de assinatura) e SICONV
    (propostas aprovadas e ainda nao assinadas)."""
    out = {}
    pp = g.carregar("df_parcerias_proposta.csv", dtype=str)
    par = g.carregar("df_parcerias_parceria.csv", dtype=str)
    m = par.merge(pp[["id_proposta", "dt_proposta"]], on="id_proposta", how="left")
    for cd, dt in zip(m["cd_parceria"].astype(str).str.replace(r"\.0$", "", regex=True), m["dt_proposta"]):
        out[("Gestão de Parcerias", g.chave_id(cd))] = pd.to_datetime(dt, errors="coerce")
    pr = g.carregar("df_siconv_proposta.csv", dtype=str)
    cv = g.carregar("df_siconv_convenio.csv", dtype=str)
    m = cv.merge(pr[["ID_PROPOSTA", "DIA_PROPOSTA"]], on="ID_PROPOSTA", how="left")
    for nr, dt in zip(m["NR_CONVENIO"], m["DIA_PROPOSTA"]):
        out[("Discricionárias e Legais", g.chave_id(nr))] = pd.to_datetime(dt, dayfirst=True, errors="coerce")
    return out


def montar_universo():
    dados = {}
    for aba, f in (("siconv", g.montar_siconv), ("parcerias", g.montar_parcerias),
                   ("especiais", g.montar_especiais), ("fundoafundo", g.montar_fundoafundo)):
        dados[aba] = f()[0]
    cz, _, cobertura = g.montar_cruzamento()

    obj_sic = objeto_siconv()
    faf_raw = g.carregar("df_fundoafundo_plano_acao.csv", dtype=str)
    obj_faf = faf_raw.drop_duplicates("codigo_plano_acao").set_index("codigo_plano_acao")["objetivos_plano_acao"]
    esp_raw = g.carregar("df_especiais_plano_acao.csv", dtype=str)
    obj_esp = esp_raw.drop_duplicates("codigo_plano_acao").set_index("codigo_plano_acao")["nome_objeto"]

    dt_prop = datas_proposta()
    linhas = []
    for fonte, info in g.FONTE_TAB_INFO.items():
        df = dados[info["aba"]]
        idcol, aba = info["id_col"], info["aba"]
        cob = cobertura.get(fonte, {})
        pendentes = {g.chave_id(x) for x in cob.get("ids_nao_localizados", set())}
        achados = {g.chave_id(k): v for k, v in cob.get("ids_localizados", {}).items()}
        achados_siafi = {g.chave_id(k): v for k, v in cob.get("nutransfsiafi_localizados", {}).items()}
        cand = {g.chave_id(k): v for k, v in cob.get("candidatos_semelhanca", {}).items()}
        for _, r in df.iterrows():
            ident = str(r[idcol]).strip()
            chave = g.chave_id(ident)
            executado = float(g.num(pd.Series([r[info["valor_col"]]])).iloc[0])
            if aba == "siconv":
                objeto, benef, emenda = obj_sic.get(ident, ""), r["Beneficiário"], r["Nº Emenda Parlamentar"]
            elif aba == "parcerias":
                objeto, benef, emenda = r["Objeto"], r["Beneficiário"], ""
            elif aba == "especiais":
                objeto, emenda = obj_esp.get(ident, ""), r["Nº Emenda Parlamentar"]
                benef = r["Órgão executor"] or r["Beneficiário"]
            else:
                objeto, emenda = obj_faf.get(ident, ""), ""
                benef = r.get("Fundo vinculado") or r["Beneficiário"]
            emenda = "" if pd.isna(emenda) else (str(emenda).replace(".0", "") if str(emenda).replace(".0", "").isdigit() else emenda)
            data = r["Data de celebração"]
            linhas.append({
                "Fonte": fonte, "Identificador": ident, "Concedente": r["Concedente"],
                "Beneficiário / Órgão executor": benef,
                "Objeto": re.sub(r"\s+", " ", str(objeto if pd.notna(objeto) else "")).strip(),
                "Data de Celebração": data,
                "Data da Proposta": dt_prop.get((fonte, chave), pd.NaT),
                "Valor Global (R$)": float(g.num(pd.Series([r["Valor global (R$)"]])).iloc[0]),
                "Valor Repassado (R$)": executado,
                "Situação": r["Situação"],
                "Nº da Emenda": emenda,
                "Execução financeira": "Com execução" if executado > 0 else "Sem execução",
                "Candidato no SIGGo (Nº Transf. / camada)": ", ".join(cand.get(chave, [])),
                "Localizado no SIGGo": "Não" if chave in pendentes else "Sim",
                "NUTRANSFSIAFI no SIGGo": ", ".join(achados_siafi.get(chave, [])),
                "Nº Transf. SIGGo": ", ".join(achados.get(chave, [])),
                "Em fase ativa": "Sim" if em_fase_ativa(r["Situação"], executado, pd.notna(data) and str(data) != "") else "Não",
            })
    todos = pd.DataFrame(linhas, columns=COLUNAS_TODOS)
    todos["_dt"] = pd.to_datetime(todos["Data de Celebração"], dayfirst=True, errors="coerce")
    todos["Data da Proposta"] = pd.to_datetime(todos["Data da Proposta"], errors="coerce")
    todos["_ord"] = todos["_dt"].fillna(todos["Data da Proposta"])  # ordenacao: celebracao, senao proposta
    pend = todos[(todos["Localizado no SIGGo"] == "Não") & (todos["Em fase ativa"] == "Sim")].copy()
    for d in (pend, todos):
        d["Data de Celebração"] = d["_dt"]
    pend = pend.sort_values("_ord", ascending=False, na_position="last")
    todos = todos.sort_values("_ord", ascending=False, na_position="last")
    return pend[COLUNAS_PEND + ["Candidato no SIGGo (Nº Transf. / camada)"]], todos.drop(columns=["_dt", "_ord"]), cobertura


LARGURAS = {"Fonte": 24, "Identificador": 24, "Concedente": 34, "Beneficiário / Órgão executor": 42, "Objeto": 70,
            "Data de Celebração": 14, "Data da Proposta": 14, "Valor Global (R$)": 18, "Valor Repassado (R$)": 18, "Situação": 30, "Nº da Emenda": 16,
            "Execução financeira": 18, "Candidato no SIGGo (Nº Transf. / camada)": 30, "Localizado no SIGGo": 14,
            "NUTRANSFSIAFI no SIGGo": 24, "Nº Transf. SIGGo": 16, "Em fase ativa": 12}


def valor_excel(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def escrever(ws, df, colunas, titulo, notas, colorir=False):
    """Cabecalho: titulo (linha 1) + duas linhas de explicacao (2 e 3, mescladas e
    com quebra de texto) + linha de colunas (4)."""
    ultima = get_column_letter(len(colunas))
    ws.append([titulo]); ws["A1"].font = Font(bold=True, size=13)
    for i, texto in enumerate(notas, start=2):
        ws.append([texto])
        ws.merge_cells(f"A{i}:{ultima}{i}")
        ws[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"A{i}"].font = Font(italic=(i == 3), bold=(i == 2), color="333333" if i == 2 else "555555")
        ws.row_dimensions[i].height = 15 * (len(texto) // 170 + 1)
    ws.append(colunas)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0D1B3E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in df[colunas].itertuples(index=False):
        ws.append([valor_excel(v) for v in row])
    idx = {c: i for i, c in enumerate(colunas)}
    for c, i in idx.items():
        ws.column_dimensions[get_column_letter(i + 1)].width = LARGURAS.get(c, 16)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for i, c in enumerate(row):
            c.alignment = Alignment(wrap_text=(colunas[i] == "Objeto"), vertical="top")
        for nome in ("Valor Global (R$)", "Valor Repassado (R$)"):
            row[idx[nome]].number_format = "#,##0.00"
        for nome in ("Data de Celebração", "Data da Proposta"):
            row[idx[nome]].number_format = "dd/mm/yyyy"
        if colorir:
            fill = VERDE if row[idx["Localizado no SIGGo"]].value == "Sim" else VERMELHO
            for c in row:
                c.fill = fill
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(colunas))}{max(ws.max_row, 5)}"


def main():
    pend, todos, cobertura = montar_universo()
    hoje = datetime.now()
    datas = g.carregar_datas_atualizacao()
    ref = "; ".join(f"{k}: {v}" for k, v in datas.items())
    n_exec = int((pend["Execução financeira"] == "Com execução").sum())

    wb = Workbook()
    ws = wb.active; ws.title = "Resumo"
    ws2 = wb.create_sheet("Pendências")
    n_nao_ident = int((todos["Localizado no SIGGo"] == "Não").sum())
    n_inativos = n_nao_ident - len(pend)
    escrever(ws2, pend, COLUNAS_PEND,
             "Instrumentos do TransfereGov (GDF) sem identificador cadastrado no SIGGo (NUTRANSFSIAFI) — em fase ativa",
             [f"O que é pendência: instrumento do TransfereGov (GDF) cujo identificador não consta no campo NUTRANSFSIAFI de nenhum registro do SIGGo "
              f"e que está em fase ativa (aprovado, em execução, autorizado, prestação de contas etc.), com ou sem execução financeira. "
              f"A execução financeira não é critério: só separa as pendências em {n_exec} com execução e {len(pend) - n_exec} sem execução. "
              f"Dos {n_nao_ident} instrumentos não identificados, {len(pend)} são pendências e {n_inativos} ficam de fora por estarem encerrados ou não formalizados "
              f"(anulado, cancelado, rescindido, impedido, rejeitado, em elaboração, ou sem situação, data e execução); esses constam na aba “Todos (APIs e CSV)” com “Em fase ativa = Não”.",
              f"Gerado em {hoje:%d/%m/%Y %H:%M}. Origem atualizada em — {ref}. O “candidato” é só uma pista (Nº Original, conta, objeto, beneficiário/valor), não identificação. "
              "Valor Repassado = valor pago/repassado/desembolsado na origem (Fundo a Fundo: conta compartilhada por vários planos é rateada pelo repasse). "
              "Data de Celebração: no Fundo a Fundo é o início de vigência; em Gestão de Parcerias nenhuma parceria do GDF traz data de assinatura. "
              "Data da Proposta: preenchida em Gestão de Parcerias e Discricionárias e Legais; a ordenação usa a celebração e, na falta dela, a data da proposta."])
    ws3 = wb.create_sheet("Todos (APIs e CSV)")
    escrever(ws3, todos, COLUNAS_TODOS,
             "Todos os instrumentos encontrados nas APIs/CSVs do TransfereGov (GDF)",
             ["Verde = identificador cadastrado no NUTRANSFSIAFI de algum registro do SIGGo (valor indicado na coluna “NUTRANSFSIAFI no SIGGo”); vermelho = não cadastrado. "
              "“Em fase ativa” indica se o instrumento vermelho entra na aba Pendências (Sim) ou fica de fora por estar encerrado ou não formalizado (Não).",
              f"Gerado em {hoje:%d/%m/%Y %H:%M}. “NUTRANSFSIAFI no SIGGo” é o valor gravado no campo; “Nº Transf. SIGGo” é o número da transferência no SIGGo."],
             colorir=True)

    ws.append(["Fonte", "Instrumentos", "Identificado no SIGGo", "Não identificado no SIGGo", "Pendências",
               "Pendências com execução financeira", "Valor Repassado – todos os instrumentos (R$)",
               "Valor Repassado – pendências (R$)", "Pendências sem execução financeira"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0D1B3E")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for fonte in g.FONTE_TAB_INFO:
        t = todos[todos["Fonte"] == fonte]; p = pend[pend["Fonte"] == fonte]
        com = p[p["Execução financeira"] == "Com execução"]
        ws.append([fonte, len(t), int((t["Localizado no SIGGo"] == "Sim").sum()), int((t["Localizado no SIGGo"] == "Não").sum()),
                   len(p), len(com), float(t["Valor Repassado (R$)"].sum()), float(com["Valor Repassado (R$)"].sum()),
                   len(p) - len(com)])
    n = ws.max_row
    ws.append(["TOTAL"] + [f"=SUM({get_column_letter(i)}2:{get_column_letter(i)}{n})" for i in range(2, 10)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in ws.iter_rows(min_row=2):
        r[6].number_format = "#,##0.00"; r[7].number_format = "#,##0.00"
    for i, w in enumerate([28, 14, 18, 20, 14, 20, 24, 24, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 45

    saida = g.PASTA / f"relatorio_pendencias_transferegov_{hoje:%Y%m%d}.xlsx"
    try:
        wb.save(saida)
    except PermissionError:  # arquivo do dia aberto no Excel: grava com sufixo de hora
        saida = saida.with_name(f"{saida.stem}_{hoje:%H%M}.xlsx")
        wb.save(saida)
    print(f"Relatorio salvo em {saida}")
    print(f"  pendencias em fase ativa: {len(pend)} (com execucao: {n_exec}) | total de instrumentos nas APIs: {len(todos)}")


if __name__ == "__main__":
    main()
